"""Builds excel/analysis.xlsx (Raw, Lookup, Clean, Summary) with live formulas."""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

ROOT = Path(__file__).resolve().parents[1]
F = "Arial"
NAVY, BLUE, LIGHT, GREY = "1F3864", "2F5597", "D9E2F3", "F2F2F2"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def hdr(c, fill=NAVY):
    c.font = Font(name=F, bold=True, color="FFFFFF", size=11)
    c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER

def body(c, fmt=None, bold=False, center=True, fill=None, color="000000"):
    c.font = Font(name=F, size=10, bold=bold, color=color)
    c.alignment = Alignment(horizontal="center" if center else "left", vertical="center")
    c.border = BORDER
    if fmt: c.number_format = fmt
    if fill: c.fill = PatternFill("solid", fgColor=fill)

def read(p):
    with open(ROOT / p, newline="", encoding="utf-8") as f:
        return list(csv.reader(f))

fact, look = read("data/raw/assessments.csv"), read("data/raw/courses.csv")
wb = Workbook()

# ------------------------------------------------------------------ Raw (13 rows, unchanged)
ws = wb.active; ws.title = "Raw"
for j, h in enumerate(fact[0], 1): hdr(ws.cell(1, j, h))
for i, r in enumerate(fact[1:], 2):
    ws.cell(i, 1, int(r[0])); ws.cell(i, 2, r[1]); ws.cell(i, 3, r[2]); ws.cell(i, 4, r[3])
    ws.cell(i, 5, float(r[4])); ws.cell(i, 6, float(r[5]))
    dup = (i == 14)
    for j in range(1, 7):
        body(ws.cell(i, j), "0" if j in (1, 5, 6) else None, fill="FCE4D6" if dup else None)
ws["H1"] = "Notes"; hdr(ws["H1"], BLUE)
ws["H2"] = "Raw count (rows)"; ws["I2"] = "=COUNTA(A2:A14)"
ws["H3"] = "Row 14 (shaded) is an exact duplicate of row 13 - removed in the Clean sheet."
for c in ("H2", "I2"): body(ws[c], bold=True, center=(c == "I2"))
ws["H3"].font = Font(name=F, size=9, italic=True, color="C00000")
for col, w in zip("ABCDEFGHI", (15, 10, 11, 12, 9, 15, 3, 18, 10)): ws.column_dimensions[col].width = w
ws.column_dimensions["H"].width = 24
ws.freeze_panes = "A2"

# ------------------------------------------------------------------ Lookup
wl = wb.create_sheet("Lookup")
for j, h in enumerate(look[0], 1): hdr(wl.cell(1, j, h))
for i, r in enumerate(look[1:], 2):
    for j, v in enumerate(r, 1): body(wl.cell(i, j, v), center=(j == 1))
for col, w in zip("ABC", (11, 14, 16)): wl.column_dimensions[col].width = w
wl.freeze_panes = "A2"

# ------------------------------------------------------------------ Clean (12 rows)
wc = wb.create_sheet("Clean")
heads = fact[0] + ["department", "pass_flag"]
for j, h in enumerate(heads, 1): hdr(wc.cell(1, j, h))
seen, rows = set(), []
for r in fact[1:]:
    t = tuple(r)
    if t not in seen: seen.add(t); rows.append(r)
assert len(rows) == 12
for i, r in enumerate(rows, 2):
    wc.cell(i, 1, int(r[0])); wc.cell(i, 2, r[1]); wc.cell(i, 3, r[2]); wc.cell(i, 4, r[3])
    wc.cell(i, 5, float(r[4])); wc.cell(i, 6, float(r[5]))
    wc.cell(i, 7, f"=INDEX(Lookup!$C$2:$C$5,MATCH(C{i},Lookup!$A$2:$A$5,0))")   # XLOOKUP-equivalent
    wc.cell(i, 8, f"=IF(E{i}>=50,1,0)")
    for j in range(1, 9):
        body(wc.cell(i, j), "0" if j in (1, 5, 6, 8) else None)
wc["J1"] = "Row-count check"; wc["K1"] = "Rows"
hdr(wc["J1"], BLUE); hdr(wc["K1"], BLUE)
wc["J2"] = "Before (Raw)";        wc["K2"] = "=COUNTA(Raw!A2:A14)"
wc["J3"] = "After (Clean)";       wc["K3"] = "=COUNTA(A2:A13)"
wc["J4"] = "Duplicates removed";  wc["K4"] = "=K2-K3"
for r in (2, 3, 4):
    body(wc.cell(r, 10), center=False, bold=True); body(wc.cell(r, 11), "0", bold=True, fill=LIGHT)
wc["J6"] = "department = INDEX/MATCH on Lookup (XLOOKUP equivalent); pass_flag = IF(score>=50,1,0)."
wc["J6"].font = Font(name=F, size=9, italic=True, color="595959")
for col, w in zip("ABCDEFGHIJK", (15, 10, 11, 12, 9, 15, 14, 11, 3, 22, 9)): wc.column_dimensions[col].width = w
wc.freeze_panes = "A2"

# ------------------------------------------------------------------ Summary
s = wb.create_sheet("Summary")
s.sheet_view.showGridLines = False
for col, w in zip("ABCDEFG", (26, 14, 14, 14, 14, 14, 14)): s.column_dimensions[col].width = w

s.merge_cells("A1:G1"); s["A1"] = "Training Performance Analysis - Summary (Set C)"
s["A1"].font = Font(name=F, size=18, bold=True, color="FFFFFF")
s["A1"].fill = PatternFill("solid", fgColor=NAVY)
s["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
s.row_dimensions[1].height = 34
s.merge_cells("A2:G2")
s["A2"] = "Business question: which course needs the most academic support, and how does performance differ across batches?  (12 clean assessments; pass = score >= 50)"
s["A2"].font = Font(name=F, size=9, italic=True, color="595959")
s["A2"].alignment = Alignment(wrap_text=True, vertical="center", indent=1)
s.row_dimensions[2].height = 26

# KPI cards (row 4 label, row 5 value)
kpis = [("A", "Assessments (clean)", "=COUNTA(Clean!A2:A13)", "0"),
        ("C", "Average Score", "=AVERAGE(Clean!E2:E13)", "0.00"),
        ("E", "Overall Pass Rate", "=SUM(Clean!H2:H13)/COUNTA(Clean!A2:A13)", "0.00%"),
        ("G", "Rows Removed", "=Clean!K4", "0")]
spans = {"A": "B", "C": "D", "E": "F", "G": "G"}
for col, label, fml, fmt in kpis:
    end = spans[col]
    if end != col:
        s.merge_cells(f"{col}4:{end}4"); s.merge_cells(f"{col}5:{end}5")
    s[f"{col}4"] = label; s[f"{col}5"] = fml
    s[f"{col}4"].font = Font(name=F, size=9, bold=True, color="FFFFFF")
    s[f"{col}4"].fill = PatternFill("solid", fgColor=BLUE)
    s[f"{col}4"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    s[f"{col}5"].font = Font(name=F, size=18, bold=True, color=NAVY)
    s[f"{col}5"].fill = PatternFill("solid", fgColor=LIGHT)
    s[f"{col}5"].alignment = Alignment(horizontal="center", vertical="center")
    s[f"{col}5"].number_format = fmt
    for c in (col, end):
        s[f"{c}4"].border = BORDER; s[f"{c}5"].border = BORDER
s.row_dimensions[4].height = 22; s.row_dimensions[5].height = 36

def title(r, text, last="G"):
    s.merge_cells(f"A{r}:{last}{r}"); s[f"A{r}"] = text
    s[f"A{r}"].font = Font(name=F, size=12, bold=True, color=NAVY)
    s[f"A{r}"].border = Border(bottom=Side(style="medium", color=NAVY))

# Table 1 - COUNTIFS by batch
title(7, "1. Passing assessments by batch (COUNTIFS)")
for j, h in enumerate(["Batch", "Passing (pass_flag=1)", "Total Assessments", "Pass Rate"], 1):
    hdr(s.cell(8, j, h))
s.row_dimensions[8].height = 30
for i, b in enumerate(["Morning", "Evening", "Weekend"], 9):
    s.cell(i, 1, b)
    s.cell(i, 2, f'=COUNTIFS(Clean!$D$2:$D$13,A{i},Clean!$H$2:$H$13,1)')
    s.cell(i, 3, f'=COUNTIFS(Clean!$D$2:$D$13,A{i})')
    s.cell(i, 4, f'=IF(C{i}=0,0,B{i}/C{i})')
    for j in range(1, 5): body(s.cell(i, j), [None, "0", "0", "0.00%"][j - 1], center=(j > 1), bold=(j == 1))
s["A12"] = "Total"; s["B12"] = "=SUM(B9:B11)"; s["C12"] = "=SUM(C9:C11)"; s["D12"] = "=IF(C12=0,0,B12/C12)"
for j in range(1, 5): body(s.cell(12, j), [None, "0", "0", "0.00%"][j - 1], bold=True, fill=LIGHT, center=(j > 1))

# Table 2 - average score dept x month (pivot-style, live formulas)
title(14, "2. Average score by department and month (PivotTable-style cross-tab, live formulas)")
for j, h in enumerate(["Department", "Jan", "Feb", "Mar", "Overall"], 1): hdr(s.cell(15, j, h))
for i, d in enumerate(["Business", "Technology"], 16):
    s.cell(i, 1, d)
    for j, m in enumerate(["B", "C", "D"], 2):
        s.cell(i, j, f'=IFERROR(AVERAGEIFS(Clean!$E$2:$E$13,Clean!$G$2:$G$13,$A{i},Clean!$B$2:$B$13,{m}$15),0)')
    s.cell(i, 5, f'=IFERROR(AVERAGEIFS(Clean!$E$2:$E$13,Clean!$G$2:$G$13,$A{i}),0)')
    for j in range(1, 6): body(s.cell(i, j), None if j == 1 else "0.00", center=(j > 1), bold=(j == 1))
s["A18"] = "All departments"
for j, m in enumerate(["B", "C", "D"], 2):
    s.cell(18, j, f'=IFERROR(AVERAGEIFS(Clean!$E$2:$E$13,Clean!$B$2:$B$13,{m}$15),0)')
s["E18"] = "=AVERAGE(Clean!E2:E13)"
for j in range(1, 6): body(s.cell(18, j), None if j == 1 else "0.00", bold=True, fill=LIGHT, center=(j > 1))

# Table 3 - course support ranking
title(20, "3. Course performance - which course needs the most support?")
for j, h in enumerate(["Course", "Department", "Total", "Passing", "Pass Rate", "Avg Score"], 1): hdr(s.cell(21, j, h))
for i, cid in enumerate(["C1", "C2", "C3", "C4"], 22):
    k = i - 20  # Lookup row (2..5)
    s.cell(i, 1, f"=Lookup!B{k}"); s.cell(i, 2, f"=Lookup!C{k}")
    s.cell(i, 3, f"=COUNTIFS(Clean!$C$2:$C$13,Lookup!A{k})")
    s.cell(i, 4, f"=COUNTIFS(Clean!$C$2:$C$13,Lookup!A{k},Clean!$H$2:$H$13,1)")
    s.cell(i, 5, f"=IF(C{i}=0,0,D{i}/C{i})")
    s.cell(i, 6, f"=IFERROR(AVERAGEIFS(Clean!$E$2:$E$13,Clean!$C$2:$C$13,Lookup!A{k}),0)")
    for j in range(1, 7):
        body(s.cell(i, j), [None, None, "0", "0", "0.00%", "0.00"][j - 1], center=(j > 2), bold=(j == 1))
s.merge_cells("A26:C26"); s["A26"] = "Lowest pass-rate course (needs most support):"
s["A26"].font = Font(name=F, size=10, bold=True, color="C00000")
s["D26"] = "=INDEX(A22:A25,MATCH(MIN(E22:E25),E22:E25,0))"
s["E26"] = "=INDEX(D22:D25,MATCH(MIN(E22:E25),E22:E25,0))&\" / \"&INDEX(C22:C25,MATCH(MIN(E22:E25),E22:E25,0))"
s["F26"] = "=MIN(E22:E25)"
for c, f in (("D26", None), ("E26", None), ("F26", "0.00%")):
    s[c].font = Font(name=F, size=11, bold=True, color="C00000")
    s[c].alignment = Alignment(horizontal="center"); 
    if f: s[c].number_format = f
s["A27"] = "If courses tie for the lowest rate, the first is shown; compare the Pass Rate column for ties."
s["A27"].font = Font(name=F, size=8, italic=True, color="7F7F7F")

# Chart from cross-tab
title(29, "4. Chart - average score by department and month")
ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"
ch.title = "Average Score by Department and Month"
ch.x_axis.title = "Department"; ch.y_axis.title = "Average Score (0-100)"
data = Reference(s, min_col=2, max_col=4, min_row=15, max_row=17)   # Jan..Mar with header row
cats = Reference(s, min_col=1, min_row=16, max_row=17)
ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 100; ch.y_axis.majorUnit = 20
ch.x_axis.delete = False; ch.y_axis.delete = False
ch.legend.position = "r"
ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True
ch.dataLabels.showSerName = False; ch.dataLabels.showCatName = False
ch.dataLabels.showLegendKey = False; ch.dataLabels.showPercent = False
ch.dataLabels.numFmt = "0.00"; ch.y_axis.number_format = "0"
ch.gapWidth = 80
for ser, colr in zip(ch.series, ("9DC3E6", "2E75B6", "1F3864")):
    ser.graphicalProperties.solidFill = colr
ch.height, ch.width = 9.0, 20.0
s.add_chart(ch, "A30")

s["A49"] = ("Notes: Cross-tab above uses AVERAGEIFS so it stays live and error-free; to add a native PivotTable "
            "select Clean!A1:H13 > Insert > PivotTable (Rows: department, Columns: month, Values: Average of score).")
s["A49"].font = Font(name=F, size=8, italic=True, color="7F7F7F")
s.merge_cells("A49:G50"); s["A49"].alignment = Alignment(wrap_text=True, vertical="top")

from openpyxl.worksheet.properties import PageSetupProperties
for sh in wb.worksheets:
    sh.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sh.page_setup.fitToWidth = 1; sh.page_setup.fitToHeight = 1 if sh.title == "Summary" else 0
s.page_setup.orientation = "portrait"
wb.active = wb.worksheets.index(s)
wb.save(ROOT / "excel/analysis.xlsx")
print("saved")
